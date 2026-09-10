import streamlit as st
import io
import re
import base64
from datetime import datetime, date

from pptx import Presentation
from PIL import Image

# ============================================================
# CONFIG
# ============================================================

st.set_page_config(page_title="Monthly Maintenance Highlight Builder", layout="wide")

TOPIC_ORDER = ["Executive Summary", "Complete Work", "Ongoing Work", "Detail Work"]

# Pattern ชื่อไฟล์: YYYY-MM_W{n}_{หน่วยงาน}.pptx เช่น 2026-08_W1_O21.pptx
FNAME_PATTERN = re.compile(r"(\d{4})-(\d{2})_W(\d+)_([A-Za-z0-9]+)", re.IGNORECASE)

MONTH_ABBR = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
              "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}
DATE_RANGE_PATTERN = re.compile(r"(\d{1,2})\s*-\s*(\d{1,2})\s*([A-Za-z]+)\s*(\d{4})")

# label ที่ใช้จับคู่กับกล่องค่าด้านล่าง (nearest shape below บนสไลด์)
LABEL_MAP = {
    "equipment": "tag",
    "background /information": "background",
    "background/information": "background",
    "possible cause": "possible_cause",
    "action": "action",
}


def parse_filename_meta(filename: str):
    m = FNAME_PATTERN.search(filename)
    if not m:
        return None
    year, month, week, unit = m.groups()
    try:
        month_name = datetime(int(year), int(month), 1).strftime("%B %Y")
    except ValueError:
        month_name = f"{year}-{month}"
    return {"year": year, "month": month, "week": week, "unit": unit.upper(), "month_name": month_name}


# ============================================================
# PASSWORD GATE
# ============================================================

def check_password():
    def password_entered():
        correct = st.secrets.get("APP_PASSWORD", "")
        if st.session_state.get("password_input", "") == correct and correct:
            st.session_state["password_correct"] = True
            del st.session_state["password_input"]
        else:
            st.session_state["password_correct"] = False

    if st.session_state.get("password_correct", False):
        return True

    st.title("🔒 Monthly Maintenance Highlight Builder")
    st.text_input("รหัสผ่าน", type="password", on_change=password_entered, key="password_input")
    if "password_correct" in st.session_state and not st.session_state["password_correct"]:
        st.error("รหัสผ่านไม่ถูกต้อง")
    return False


if not check_password():
    st.stop()

# ============================================================
# HELPERS — อ่านโครงสร้างการ์ดงานจากสไลด์ PPTX
# ============================================================

ROW_TOLERANCE = 150000  # EMU (~0.16 inch) — ถือว่าอยู่แถวเดียวกันถ้าต่างกันไม่เกินนี้


def extract_work_item(slide):
    """
    อ่าน 1 สไลด์ = 1 การ์ดงาน (Tag/Plant/Date/Background/Possible cause/Action)
    จับคู่ label -> value โดย: ถ้ามีกล่องอยู่แถวเดียวกัน (Y ใกล้กัน) ทางขวา ใช้ก่อน (เช่น Equipment | B-F-110)
    ถ้าไม่มี ใช้กล่องที่อยู่ด้านล่างใกล้ที่สุดแทน (เช่น Action label แล้วเนื้อหาอยู่ใต้)
    """
    items = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            items.append({"top": shape.top, "left": shape.left, "text": shape.text_frame.text.strip()})

    fields = {"plant": "", "date_raw": "", "tag": "", "background": "", "possible_cause": "", "action": ""}

    remaining = []
    for it in items:
        low = it["text"].lower()
        if low.startswith("plant"):
            fields["plant"] = it["text"][5:].strip()
            continue
        if low.startswith("date"):
            val = it["text"][4:].strip()
            if val.startswith(":"):
                val = val[1:].strip()
            fields["date_raw"] = val
            continue
        remaining.append(it)

    used_idx = set()
    for i, it in enumerate(remaining):
        norm = it["text"].strip().lower()
        if norm not in LABEL_MAP:
            continue
        target_key = LABEL_MAP[norm]

        # 1) หากล่องแถวเดียวกันทางขวาก่อน
        best_j, best_dist = None, None
        for j, it2 in enumerate(remaining):
            if j == i or j in used_idx:
                continue
            if it2["text"].strip().lower() in LABEL_MAP:
                continue
            if abs(it2["top"] - it["top"]) <= ROW_TOLERANCE and it2["left"] > it["left"]:
                dist = it2["left"] - it["left"]
                if best_dist is None or dist < best_dist:
                    best_dist, best_j = dist, j

        # 2) ถ้าไม่เจอ ใช้กล่องด้านล่างที่ใกล้ที่สุดแทน
        if best_j is None:
            for j, it2 in enumerate(remaining):
                if j == i or j in used_idx:
                    continue
                if it2["text"].strip().lower() in LABEL_MAP:
                    continue
                if it2["top"] > it["top"]:
                    dist = it2["top"] - it["top"]
                    if best_dist is None or dist < best_dist:
                        best_dist, best_j = dist, j

        if best_j is not None:
            fields[target_key] = remaining[best_j]["text"].strip()
            used_idx.add(best_j)

    return fields


def parse_date_range(raw: str):
    if not raw:
        return None, None
    m = DATE_RANGE_PATTERN.search(raw)
    if not m:
        return None, None
    d1, d2, mon, year = m.groups()
    month_num = MONTH_ABBR.get(mon[:3].lower())
    if not month_num:
        return None, None
    try:
        return date(int(year), month_num, int(d1)), date(int(year), month_num, int(d2))
    except ValueError:
        return None, None


def format_date_range(start: date, end: date) -> str:
    if start.year != end.year:
        return f"{start.strftime('%-d %b %Y')} - {end.strftime('%-d %b %Y')}"
    if start.month != end.month:
        return f"{start.strftime('%-d %b')} - {end.strftime('%-d %b %Y')}"
    return f"{start.day} - {end.day} {start.strftime('%b %Y')}"


def normalize_tag(tag: str) -> str:
    return re.sub(r"\s+", "", (tag or "").strip().upper())


def split_lines(text: str):
    return [l.strip() for l in (text or "").splitlines() if l.strip()]


def dedup_lines(lines):
    seen, result = set(), []
    for l in lines:
        key = l.lower()
        if key not in seen:
            seen.add(key)
            result.append(l)
    return result


def parse_pptx_files(uploaded_files, week_labels):
    """คืนค่า list ของงานแต่ละการ์ด (ยังไม่รวม Tag ซ้ำ)"""
    raw_items = []
    unmatched = []
    for f, week_label in zip(uploaded_files, week_labels):
        prs = Presentation(io.BytesIO(f.getvalue()))
        for slide in prs.slides:
            fields = extract_work_item(slide)
            if not fields["tag"]:
                # เก็บไว้ให้ตรวจสอบ ไม่ทิ้งข้อมูล
                all_text = "\n".join(
                    shape.text_frame.text.strip()
                    for shape in slide.shapes
                    if shape.has_text_frame and shape.text_frame.text.strip()
                )
                if all_text:
                    unmatched.append({"week": week_label, "file": f.name, "text": all_text})
                continue

            start, end = parse_date_range(fields["date_raw"])
            problem = "\n".join([fields["background"], fields["possible_cause"]]).strip()
            raw_items.append({
                "tag": fields["tag"],
                "plant": fields["plant"],
                "problem": problem,
                "action": fields["action"],
                "date_raw": fields["date_raw"],
                "date_start": start,
                "date_end": end,
                "week_label": week_label,
            })
    return raw_items, unmatched


def merge_items_by_tag(items):
    """รวมงาน Tag เดียวกันจากหลายสัปดาห์ให้เหลือรายการเดียว"""
    groups, order = {}, []
    for it in items:
        key = normalize_tag(it["tag"])
        if not key:
            continue
        if key not in groups:
            groups[key] = {
                "display_tag": it["tag"], "plants": [], "problem_lines": [],
                "action_lines": [], "starts": [], "ends": [], "raw_dates": [], "weeks": [],
            }
            order.append(key)
        g = groups[key]
        if it["plant"] and it["plant"] not in g["plants"]:
            g["plants"].append(it["plant"])
        g["problem_lines"].extend(split_lines(it["problem"]))
        g["action_lines"].extend(split_lines(it["action"]))
        if it["date_start"]:
            g["starts"].append(it["date_start"])
        if it["date_end"]:
            g["ends"].append(it["date_end"])
        if it["date_raw"] and it["date_raw"] not in g["raw_dates"]:
            g["raw_dates"].append(it["date_raw"])
        if it["week_label"] not in g["weeks"]:
            g["weeks"].append(it["week_label"])

    merged = []
    for key in order:
        g = groups[key]
        if g["starts"] and g["ends"]:
            date_text = format_date_range(min(g["starts"]), max(g["ends"]))
        else:
            date_text = " / ".join(g["raw_dates"])
        merged.append({
            "key": key,
            "tag": g["display_tag"],
            "plant": ", ".join(g["plants"]),
            "problem": "\n".join(dedup_lines(g["problem_lines"])),
            "action": "\n".join(dedup_lines(g["action_lines"])),
            "date_range": date_text,
            "weeks": ", ".join(g["weeks"]),
        })
    return merged


# ============================================================
# HELPERS — ภาพ
# ============================================================

def resize_image_b64(blob: bytes, max_side=1000, quality=80):
    try:
        im = Image.open(io.BytesIO(blob)).convert("RGB")
    except Exception:
        return None
    im.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=quality)
    return base64.b64encode(buf.getvalue()).decode()


# ============================================================
# HTML REPORT TEMPLATE
# ============================================================

def card_html(item, images_b64):
    img_html = ""
    if images_b64:
        img_html = '<div class="img-grid">' + "".join(
            f'<img src="data:image/jpeg;base64,{b}" />' for b in images_b64
        ) + '</div>'
    problem_html = "".join(f'<li contenteditable="true">{l}</li>' for l in split_lines(item["problem"])) or '<li contenteditable="true">-</li>'
    action_html = "".join(f'<li contenteditable="true">{l}</li>' for l in split_lines(item["action"])) or '<li contenteditable="true">-</li>'
    return f'''
    <div class="work-card">
      <div class="card-head">
        <span class="card-tag" contenteditable="true">{item["tag"]}</span>
        <span class="card-plant" contenteditable="true">{item["plant"]}</span>
        <span class="card-date" contenteditable="true">{item["date_range"]}</span>
      </div>
      <div class="card-field"><div class="field-label">Problem</div><ul>{problem_html}</ul></div>
      <div class="card-field"><div class="field-label">Action</div><ul>{action_html}</ul></div>
      {img_html}
    </div>
'''


def build_html_report(title: str, period: str, exec_text: str, complete_items, ongoing_items, detail_items, images_by_tag: dict) -> str:
    stat_items = (
        f'<div class="meta-item"><b>{len(detail_items)}</b>อุปกรณ์ทั้งหมด</div>'
        f'<div class="meta-item"><b>{len(complete_items)}</b>งานเสร็จสิ้น</div>'
        f'<div class="meta-item"><b>{len(ongoing_items)}</b>งานที่ยังดำเนินการ</div>'
    )

    exec_html = "".join(f'<p contenteditable="true">{l}</p>' for l in exec_text.splitlines() if l.strip())

    def cards_block(items):
        return "".join(card_html(it, images_by_tag.get(it["key"], [])) for it in items)

    body = f'''
  <section>
    <div class="sec-head"><span class="sec-num">01</span><h2 contenteditable="true">Executive Summary</h2></div>
    {exec_html}
  </section>
  <section>
    <div class="sec-head"><span class="sec-num">02</span><h2 contenteditable="true">Complete Work</h2></div>
    {cards_block(complete_items) if complete_items else '<p class="empty-note">ไม่มีงานที่เสร็จสิ้นในเดือนนี้</p>'}
  </section>
  <section>
    <div class="sec-head"><span class="sec-num">03</span><h2 contenteditable="true">Ongoing Work</h2></div>
    {cards_block(ongoing_items) if ongoing_items else '<p class="empty-note">ไม่มีงานที่ยังดำเนินการอยู่</p>'}
  </section>
  <section>
    <div class="sec-head"><span class="sec-num">04</span><h2 contenteditable="true">Detail Work</h2></div>
    {cards_block(detail_items)}
  </section>
'''

    return f"""<!DOCTYPE html>
<html lang="th">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Archivo:wght@500;600;700;800&family=IBM+Plex+Sans:wght@400;500;600&family=IBM+Plex+Mono:wght@500&display=swap" rel="stylesheet">
<style>
  :root{{
    --paper:#EEEBE2; --panel:#F8F6F0; --ink:#1B2430; --ink-soft:#4A5568;
    --rule:#D8D3C4; --rust:#BF5B23; --teal:#35606B; --green:#4B7255;
  }}
  *{{box-sizing:border-box;}}
  body{{margin:0;background:var(--paper);color:var(--ink);font-family:'IBM Plex Sans',sans-serif;line-height:1.55;}}
  .sheet{{max-width:920px;margin:0 auto;padding:0 24px 80px;}}
  header{{padding:56px 0 32px;border-bottom:3px solid var(--ink);margin-bottom:40px;}}
  .plant-tag{{font-family:'IBM Plex Mono',monospace;font-size:13px;color:var(--rust);}}
  h1{{font-family:'Archivo',sans-serif;font-weight:800;font-size:clamp(28px,5vw,42px);letter-spacing:-0.01em;margin:10px 0 6px;}}
  .subdate{{font-size:16px;color:var(--ink-soft);}}
  .header-meta{{display:flex;gap:26px;margin-top:24px;flex-wrap:wrap;}}
  .meta-item{{font-size:13px;color:var(--ink-soft);}}
  .meta-item b{{display:block;font-family:'Archivo',sans-serif;font-size:20px;font-weight:700;color:var(--ink);}}
  section{{margin-bottom:48px;}}
  .sec-head{{display:flex;align-items:baseline;gap:14px;margin-bottom:18px;border-bottom:1px solid var(--rule);padding-bottom:10px;}}
  .sec-num{{font-family:'IBM Plex Mono',monospace;font-size:13px;color:var(--rust);min-width:22px;}}
  h2{{font-family:'Archivo',sans-serif;font-weight:700;font-size:22px;margin:0;outline:none;}}
  p{{margin:0 0 10px;font-size:15px;outline:none;}}
  .empty-note{{color:var(--ink-soft);font-size:14px;font-style:italic;}}
  .work-card{{border:1px solid var(--rule);border-radius:2px;padding:18px 20px;margin-bottom:14px;background:var(--panel);}}
  .card-head{{display:flex;gap:16px;flex-wrap:wrap;align-items:baseline;margin-bottom:12px;padding-bottom:10px;border-bottom:1px dashed var(--rule);}}
  .card-tag{{font-family:'Archivo',sans-serif;font-weight:700;font-size:17px;outline:none;}}
  .card-plant{{font-family:'IBM Plex Mono',monospace;font-size:12px;color:var(--teal);outline:none;}}
  .card-date{{font-family:'IBM Plex Mono',monospace;font-size:12px;color:var(--rust);margin-left:auto;outline:none;}}
  .card-field{{margin-bottom:10px;}}
  .field-label{{font-family:'IBM Plex Mono',monospace;font-size:11px;color:var(--ink-soft);text-transform:uppercase;margin-bottom:4px;}}
  ul{{margin:0;padding-left:20px;}}
  li{{margin-bottom:4px;font-size:14px;outline:none;}}
  li::marker{{color:var(--teal);}}
  .img-grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:8px;margin-top:10px;}}
  .img-grid img{{width:100%;border-radius:2px;border:1px solid var(--rule);display:block;}}
  footer{{margin-top:60px;padding-top:20px;border-top:1px solid var(--rule);font-size:12px;color:var(--ink-soft);font-family:'IBM Plex Mono',monospace;}}
  [contenteditable="true"]:hover{{background:rgba(53,96,107,0.06);cursor:text;}}
  .edit-hint{{font-family:'IBM Plex Mono',monospace;font-size:11px;color:var(--ink-soft);background:var(--panel);border:1px dashed var(--rule);padding:10px 14px;margin-bottom:32px;}}
</style>
</head>
<body>
<div class="sheet">
  <header>
    <div class="plant-tag">Monthly Maintenance Highlight</div>
    <h1 contenteditable="true">{title}</h1>
    <div class="subdate">{period}</div>
    <div class="header-meta">{stat_items}</div>
  </header>
  <div class="edit-hint">แก้ไขข้อความในหน้านี้ได้โดยตรง (คลิกแล้วพิมพ์) — เมื่อแก้เสร็จ ใช้เมนูเบราว์เซอร์ Save Page As (Webpage, HTML only) เพื่อบันทึกเวอร์ชันที่แก้ไขแล้ว</div>
{body}
  <footer>{title.upper()} — {period}</footer>
</div>
</body>
</html>"""


# ============================================================
# STREAMLIT UI
# ============================================================

st.title("📋 Monthly Maintenance Highlight Builder")
st.caption("อัปโหลด Weekly Report (PowerPoint) หลายไฟล์ → รวมงานตาม Tag No. ข้ามสัปดาห์ → เลือกสถานะ Complete/Ongoing → ได้ Monthly Report เป็น HTML")

uploaded_files = st.file_uploader(
    "อัปโหลดไฟล์ Weekly Report (.pptx) — เลือกได้หลายไฟล์ — ตั้งชื่อไฟล์ตาม pattern YYYY-MM_W{n}_{หน่วยงาน} "
    "เช่น 2026-08_W1_O21.pptx จะช่วยตั้งชื่อสัปดาห์และช่วงเวลาให้อัตโนมัติ",
    type=["pptx"], accept_multiple_files=True
)

if uploaded_files:
    st.subheader("🏷️ ตั้งชื่อสัปดาห์ให้แต่ละไฟล์")
    st.caption("ดึงจากชื่อไฟล์ให้อัตโนมัติถ้าตรง pattern — แก้ไขเองได้ถ้าต้องการ")
    week_labels = []
    cols = st.columns(min(len(uploaded_files), 4))
    for i, f in enumerate(uploaded_files):
        meta = parse_filename_meta(f.name)
        default_label = f"Week {meta['week']} ({meta['unit']})" if meta else f"Week {i+1}"
        with cols[i % len(cols)]:
            label = st.text_input(f"📄 {f.name}", value=default_label, key=f"week_label_{f.name}_{i}")
        week_labels.append(label)

    if st.button("🔍 อ่านไฟล์และรวมงานตาม Tag", type="primary"):
        with st.spinner("กำลังอ่านไฟล์ PowerPoint..."):
            raw_items, unmatched = parse_pptx_files(uploaded_files, week_labels)
            st.session_state["merged_items"] = merge_items_by_tag(raw_items)
            st.session_state["unmatched"] = unmatched

        for f in uploaded_files:
            meta = parse_filename_meta(f.name)
            if meta:
                st.session_state["report_period_input"] = meta["month_name"]
                break
        st.rerun()

with st.sidebar:
    st.header("⚙️ ตั้งค่า")
    report_title = st.text_input("ชื่อรายงาน", value="Monthly Maintenance Highlight")
    report_period = st.text_input(
        "ช่วงเวลา (เช่น September 2026)",
        value=st.session_state.get("report_period_input", datetime.now().strftime("%B %Y")),
        key="report_period_input"
    )
    if st.button("🚪 ออกจากระบบ (ล้างรหัสผ่าน)"):
        st.session_state["password_correct"] = False
        st.rerun()

if "merged_items" in st.session_state:
    merged_items = st.session_state["merged_items"]
    unmatched = st.session_state.get("unmatched", [])

    st.divider()
    st.success(f"รวมงานได้ {len(merged_items)} รายการ (นับตาม Tag No. ที่ไม่ซ้ำกัน)")

    if unmatched:
        with st.expander(f"⚠️ พบ {len(unmatched)} สไลด์ที่จับ Tag ไม่ได้ (ตรวจสอบด้วยตนเอง)"):
            for u in unmatched:
                st.caption(f"📄 {u['file']} — {u['week']}")
                st.text(u["text"])
                st.divider()

    st.subheader("✏️ ตรวจทานแต่ละ Tag — เลือกสถานะ Complete / Ongoing และแนบภาพ")

    if "item_overrides" not in st.session_state:
        st.session_state["item_overrides"] = {}
    if "item_images" not in st.session_state:
        st.session_state["item_images"] = {}
    if "item_status" not in st.session_state:
        st.session_state["item_status"] = {}

    final_items = []
    images_by_tag = {}

    for item in merged_items:
        key = item["key"]
        with st.expander(f"🔧 {item['tag']}  —  {item['plant']}  —  {item['date_range']}  (รวมจาก: {item['weeks']})", expanded=False):
            ov = st.session_state["item_overrides"].setdefault(key, {
                "plant": item["plant"], "problem": item["problem"],
                "action": item["action"], "date_range": item["date_range"],
            })

            c1, c2 = st.columns(2)
            with c1:
                ov["plant"] = st.text_input("Plant", value=ov["plant"], key=f"plant_{key}")
            with c2:
                ov["date_range"] = st.text_input("ช่วงวันที่", value=ov["date_range"], key=f"date_{key}")

            ov["problem"] = st.text_area("Problem", value=ov["problem"], height=100, key=f"problem_{key}")
            ov["action"] = st.text_area("Action", value=ov["action"], height=100, key=f"action_{key}")

            status = st.radio("สถานะงาน", ["Ongoing", "Complete"],
                               index=0 if st.session_state["item_status"].get(key, "Ongoing") == "Ongoing" else 1,
                               key=f"status_{key}", horizontal=True)
            st.session_state["item_status"][key] = status

            st.markdown("**🖼️ ภาพประกอบ**")
            count_key = f"imgcount_{key}"
            num_images = st.number_input("จำนวนกรอบภาพ", min_value=0, max_value=10,
                                          value=st.session_state.get(count_key, 0), step=1, key=count_key)
            imgs = []
            if num_images > 0:
                img_cols = st.columns(4)
                for i in range(int(num_images)):
                    with img_cols[i % 4]:
                        up = st.file_uploader(f"ภาพที่ {i+1}", type=["png", "jpg", "jpeg"], key=f"upimg_{key}_{i}")
                        if up is not None:
                            b64 = resize_image_b64(up.getvalue())
                            if b64:
                                st.image(up, use_container_width=True)
                                imgs.append(b64)
            images_by_tag[key] = imgs

            final_items.append({
                "key": key, "tag": item["tag"], "plant": ov["plant"],
                "problem": ov["problem"], "action": ov["action"],
                "date_range": ov["date_range"], "status": status,
            })

    complete_items = [it for it in final_items if it["status"] == "Complete"]
    ongoing_items = [it for it in final_items if it["status"] == "Ongoing"]

    st.divider()
    st.subheader("📝 Executive Summary")
    default_exec = (
        f"เดือนนี้มีงานทั้งหมด {len(final_items)} รายการ ครอบคลุม "
        f"{len(set(it['plant'] for it in final_items if it['plant']))} plant "
        f"— เสร็จสิ้น {len(complete_items)} รายการ, ยังดำเนินการ {len(ongoing_items)} รายการ"
    )
    if "exec_text" not in st.session_state:
        st.session_state["exec_text"] = default_exec
    exec_text = st.text_area("แก้ไขสรุปผู้บริหารได้อิสระ", value=st.session_state["exec_text"], height=100, key="exec_text_area")
    st.session_state["exec_text"] = exec_text

    st.divider()
    st.subheader("📄 Monthly Report")

    html_report = build_html_report(report_title, report_period, exec_text, complete_items, ongoing_items, final_items, images_by_tag)

    tab1, tab2 = st.tabs(["👁️ ดูตัวอย่างในแอป", "⬇️ ดาวน์โหลด"])
    with tab1:
        st.components.v1.html(html_report, height=1000, scrolling=True)
    with tab2:
        st.download_button(
            "⬇️ ดาวน์โหลดเป็นไฟล์ HTML",
            data=html_report.encode("utf-8"),
            file_name=f"Monthly_Maintenance_Highlight_{report_period.replace(' ', '_')}.html",
            mime="text/html",
        )
        st.caption("ไฟล์ที่ดาวน์โหลดสามารถเปิดด้วยเบราว์เซอร์แล้วแก้ไขข้อความต่อได้โดยตรง (คลิกแล้วพิมพ์)")
else:
    st.info("อัปโหลดไฟล์ PowerPoint รายสัปดาห์ด้านบน แล้วกด 'อ่านไฟล์และรวมงานตาม Tag' เพื่อเริ่มต้น")
